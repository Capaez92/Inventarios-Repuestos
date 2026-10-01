import io
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from app.models.models import Product, InventoryMovement, CompanySetting
from datetime import datetime

def generate_inventory_excel():
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario Valorizado"

    # Styling
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    alert_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    alert_font = Font(name="Segoe UI", size=10, color="991B1B", bold=True)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    company = CompanySetting.query.first()
    company_name = company.company_name if company else "Almacén de Repuestos"

    ws.merge_cells("A1:J1")
    ws["A1"] = company_name.upper()
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:J2")
    ws["A2"] = f"Reporte de Inventario Valorizado - Generado el: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = [
        "Código/SKU", "Código Barras", "Descripción", "Categoría", 
        "Marca", "Ubicación", "Stock", "Stock Mín", "Costo Prom.", "Valor Total Costo"
    ]
    
    ws.append([]) # row 3 blank
    ws.append(headers) # row 4
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    products = Product.query.order_by(Product.name).all()
    total_valuation = 0.0
    total_items = 0

    for idx, p in enumerate(products, start=5):
        stock = p.current_stock or 0
        min_s = p.min_stock or 0
        cost = p.cost_price or 0.0
        val = stock * cost
        total_valuation += val
        total_items += stock

        cat_name = p.category.name if p.category else 'N/A'
        brand_name = p.brand.name if p.brand else 'N/A'
        loc_name = p.location.full_location if p.location else 'Sin Asignar'

        ws.append([
            p.sku,
            p.barcode or '',
            p.name,
            cat_name,
            brand_name,
            loc_name,
            stock,
            min_s,
            cost,
            val
        ])

        for c_idx in range(1, 11):
            c = ws.cell(row=idx, column=c_idx)
            c.font = regular_font
            c.border = thin_border
            if c_idx in [7, 8]:
                c.alignment = Alignment(horizontal="center")
            elif c_idx in [9, 10]:
                c.alignment = Alignment(horizontal="right")
                c.number_format = "$#,##0.00"

        # Highlight low stock
        if stock <= min_s:
            ws.cell(row=idx, column=7).fill = alert_fill
            ws.cell(row=idx, column=7).font = alert_font

    # Totals row
    summary_row = len(products) + 5
    ws.cell(row=summary_row, column=3, value="TOTAL GENERAL").font = bold_font
    ws.cell(row=summary_row, column=7, value=total_items).font = bold_font
    ws.cell(row=summary_row, column=10, value=total_valuation).font = bold_font
    ws.cell(row=summary_row, column=10).number_format = "$#,##0.00"

    # Column widths
    widths = [15, 16, 35, 18, 16, 25, 10, 10, 14, 18]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_kardex_excel(product_id):
    product = Product.query.get_or_404(product_id)
    wb = Workbook()
    ws = wb.active
    ws.title = f"Kárdex - {product.sku}"

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E293B")
    regular_font = Font(name="Segoe UI", size=10)
    
    ws.merge_cells("A1:H1")
    ws["A1"] = f"KÁRDEX DE ARTÍCULO: {product.name} (SKU: {product.sku})"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = f"Stock Actual: {product.current_stock} {product.unit_measure} | Costo CPP: ${product.cost_price:,.2f}"
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["Fecha", "Tipo Movimiento", "Doc. Referencia", "Cantidad", "Stock Anterior", "Nuevo Stock", "Costo Unitario", "Notas"]
    ws.append([])
    ws.append(headers)

    for col in range(1, 9):
        c = ws.cell(row=4, column=col)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center")

    movements = product.movements.order_by(InventoryMovement.created_at.desc()).all()
    for row_idx, m in enumerate(movements, start=5):
        ws.append([
            m.created_at.strftime('%Y-%m-%d %H:%M'),
            m.movement_type.replace('_', ' ').title(),
            m.reference_doc or '-',
            m.quantity,
            m.previous_stock,
            m.new_stock,
            m.unit_cost,
            m.notes or ''
        ])
        for c_idx in range(1, 9):
            ws.cell(row=row_idx, column=c_idx).font = regular_font

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_inventory_pdf():
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), leftMargin=30, rightMargin=30, topMargin=30, bottomMargin=30)
    elements = []

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor('#1E293B'),
        alignment=1 # Center
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#64748B'),
        alignment=1
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10
    )

    company = CompanySetting.query.first()
    company_name = company.company_name if company else "Almacén de Repuestos"

    elements.append(Paragraph(company_name.upper(), title_style))
    elements.append(Paragraph(f"INFORME DE INVENTARIO VALORIZADO — {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style))
    elements.append(Spacer(1, 15))

    data = [["SKU", "Descripción", "Categoría", "Marca", "Ubicación", "Stock", "Mín", "Costo", "Total"]]
    
    products = Product.query.order_by(Product.name).all()
    total_val = 0.0
    for p in products:
        c_val = (p.current_stock or 0) * (p.cost_price or 0.0)
        total_val += c_val
        data.append([
            p.sku,
            Paragraph(p.name[:35], cell_style),
            p.category.name if p.category else '-',
            p.brand.name if p.brand else '-',
            p.location.name if p.location else '-',
            str(p.current_stock or 0),
            str(p.min_stock or 0),
            f"${p.cost_price:,.2f}",
            f"${c_val:,.2f}"
        ])

    data.append(["", "TOTAL VALORIZADO", "", "", "", "", "", "", f"${total_val:,.2f}"])

    col_widths = [65, 200, 90, 80, 90, 45, 45, 65, 75]
    t = Table(data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
        ('ALIGN', (5, 1), (8, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.HexColor('#FFFFFF'), colors.HexColor('#F8FAFC')]),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E2E8F0')),
    ]))

    elements.append(t)
    doc.build(elements)
    buffer.seek(0)
    return buffer
